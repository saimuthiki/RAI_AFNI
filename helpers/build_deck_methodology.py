# -*- coding: utf-8 -*-
"""
Renders the 7 per-tenet methodology slides: every repository that contributes a
check to a tenet, the mechanism behind that check, its cost and latency class,
and where it belongs in AFNI's free-first cascade.

Reads data/tenet_methodology_data.json (produced by
build_tenet_methodology_data.py). Layout is sized against the qa_matrix.py
overflow estimator so the tables pass QA by construction.
"""
import json
import os

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from build_pptx import (
    blank_slide, add_text, add_header, add_rounded,
    NAVY, TEAL, CARD_BG, WHITE, TEXT_DARK, TEXT_MUTED, TEXT_SOFT_ON_NAVY,
    AMBER, GREEN, RED_SOFT, LINE_GREY, TENET_COLORS, TENET_ORDER,
    FONT_HEAD, FONT_BODY,
)

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_DATA_PATH = os.path.join(_ROOT, "data", "tenet_methodology_data.json")

with open(_DATA_PATH, encoding="utf-8") as f:
    METHODOLOGY = json.load(f)

# --------------------------------------------------------------- LAYOUT ----
TABLE_X = 0.55
TABLE_W = 12.2
TABLE_TOP = 1.72
HEADER_H = 0.30
BODY_H_MAX = 0.30
TABLE_BOTTOM_LIMIT = 6.62   # leave room for legend + cloud band above the footer

HEADERS = ["Repository", "Mechanism", "What it does for this tenet",
           "Cost", "Latency", "Runs", "Target", "Stage"]
WEIGHTS = [1.50, 1.90, 2.60, 1.10, 1.10, 0.90, 1.00, 0.70]

HEADER_PT = 7.5
BODY_PT = 6.0
REPO_PT = 6.4

STAGE_COLORS = {
    "Stage 1": GREEN,
    "Stage 2": AMBER,
    "Stage 3": RED_SOFT,
    "Delegates": NAVY,
    "Offline": TEXT_MUTED,
}


def _col_widths():
    total = sum(WEIGHTS)
    return [TABLE_W * w / total for w in WEIGHTS]


def _set_cell(cell, lines, fill=None, align=PP_ALIGN.LEFT):
    """lines is a list of (text, size_pt, color, bold, italic) tuples."""
    cell.margin_left = Inches(0.035)
    cell.margin_right = Inches(0.035)
    cell.margin_top = Inches(0.012)
    cell.margin_bottom = Inches(0.012)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill if fill is not None else WHITE

    tf = cell.text_frame
    tf.word_wrap = True
    for i, (text, size, color, bold, italic) in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = 0.92
        run = para.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = FONT_BODY


def add_methodology_slide(prs, tenet, idx, total, finish_slide_fn):
    data = METHODOLOGY[tenet]
    rows = data["rows"]
    color = TENET_COLORS.get(tenet, TEAL)

    slide = blank_slide(prs)
    add_header(slide, f"Tenet Methodology {idx} of {total}", tenet, accent=color)
    add_text(slide, TABLE_X, 1.38, TABLE_W, 0.26,
             f"All {len(rows)} reviewed tools that contribute a check to this tenet - what the check "
             "actually is, what it costs, and which cascade stage it belongs in.",
             size=9.5, color=TEXT_MUTED, font=FONT_BODY)

    n_rows = len(rows) + 1
    body_h = min(BODY_H_MAX, (TABLE_BOTTOM_LIMIT - TABLE_TOP - HEADER_H) / max(len(rows), 1))
    table_h = HEADER_H + body_h * len(rows)

    shape = slide.shapes.add_table(n_rows, len(HEADERS), Inches(TABLE_X), Inches(TABLE_TOP),
                                   Inches(TABLE_W), Inches(table_h))
    table = shape.table
    table.first_row = False
    table.horz_banding = False

    for i, cw in enumerate(_col_widths()):
        table.columns[i].width = Inches(cw)
    table.rows[0].height = Inches(HEADER_H)
    for i in range(1, n_rows):
        table.rows[i].height = Inches(body_h)

    # header row
    for c, head in enumerate(HEADERS):
        _set_cell(table.cell(0, c), [(head, HEADER_PT, WHITE, True, False)],
                  fill=NAVY, align=PP_ALIGN.CENTER if c >= 3 else PP_ALIGN.LEFT)

    # body rows
    for r, row in enumerate(rows, start=1):
        stripe = RGBColor(0xF4, 0xF7, 0xFA) if r % 2 == 0 else WHITE
        _set_cell(table.cell(r, 0), [(row["repo"], REPO_PT, NAVY, True, False)], fill=stripe)
        _set_cell(table.cell(r, 1), [(row["mechanism"], BODY_PT, TEXT_DARK, False, False)], fill=stripe)
        _set_cell(table.cell(r, 2), [(row["functionality"], BODY_PT, TEXT_DARK, False, False)], fill=stripe)
        _set_cell(table.cell(r, 3), [(row["cost"], BODY_PT, TEXT_DARK, False, False)],
                  fill=stripe, align=PP_ALIGN.CENTER)
        _set_cell(table.cell(r, 4), [(row["latency"], BODY_PT, TEXT_DARK, False, False)],
                  fill=stripe, align=PP_ALIGN.CENTER)
        _set_cell(table.cell(r, 5), [(row["locality"], BODY_PT, TEXT_DARK, False, False)],
                  fill=stripe, align=PP_ALIGN.CENTER)
        _set_cell(table.cell(r, 6), [(row["target"], BODY_PT, TEXT_DARK, False, False)],
                  fill=stripe, align=PP_ALIGN.CENTER)
        stage = row["stage"]
        _set_cell(table.cell(r, 7), [(stage, BODY_PT, STAGE_COLORS.get(stage, TEXT_DARK), True, False)],
                  fill=stripe, align=PP_ALIGN.CENTER)

    # ------------------------------------------------------ below the table ----
    y = TABLE_TOP + table_h + 0.09
    add_text(slide, TABLE_X, y, TABLE_W, 0.20,
             "Stage 1  free + deterministic, runs on every request      "
             "Stage 2  local model, or cloud second opinion on borderline input only      "
             "Stage 3  paid / LLM-judge / heavy      Delegates  contract or taxonomy only, "
             "no detector of its own      Offline  CI and red-team only",
             size=7.2, color=TEXT_DARK, font=FONT_BODY)
    add_text(slide, TABLE_X, y + 0.19, TABLE_W, 0.18,
             "Stage = the earliest point this tool can contribute, taken from its cheapest mechanism. "
             "Latency is a range across its mechanisms, estimated from what the source actually does - "
             "not a benchmarked measurement.",
             size=6.8, color=TEXT_MUTED, italic=True, font=FONT_BODY)

    cloud = data.get("cloud_options", [])
    if cloud:
        band_y = y + 0.40
        band_h = 6.62 + 0.52 - band_y if band_y < 6.9 else 0.42
        band_h = max(0.36, min(band_h, 0.52))
        add_rounded(slide, TABLE_X, band_y, TABLE_W, band_h, NAVY, radius=0.04)
        add_text(slide, TABLE_X + 0.14, band_y + 0.04, 2.5, 0.16,
                 "STAGE 3 - CLOUD / PAID FALLBACKS", size=6.8, color=TEAL, bold=True, font=FONT_HEAD)
        add_text(slide, TABLE_X + 0.14, band_y + 0.19, TABLE_W - 0.28, band_h - 0.22,
                 "   ".join(cloud), size=6.6, color=TEXT_SOFT_ON_NAVY, font=FONT_BODY,
                 line_spacing=1.08)

    finish_slide_fn(slide, f"Methodology - {tenet}")


def add_all_methodology_slides(prs, finish_slide_fn):
    total = len(TENET_ORDER)
    for i, tenet in enumerate(TENET_ORDER, start=1):
        add_methodology_slide(prs, tenet, i, total, finish_slide_fn)
