# -*- coding: utf-8 -*-
"""Renders the 7 capability-matrix slides as native PPTX tables."""
import json
import os
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from build_pptx import (
    blank_slide, add_text, add_header, TENET_COLORS, TENET_ORDER,
    NAVY, TEAL, TEXT_DARK, TEXT_MUTED, LINE_GREY, WHITE, CARD_BG,
    FONT_HEAD, FONT_BODY,
)
from build_capability_matrix_data import DISPLAY

_DATA_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data")
with open(os.path.join(_DATA_DIR, "capability_matrix_data.json"), encoding="utf-8") as f:
    MATRIX = json.load(f)

SYM_FILL = {
    "●": RGBColor(0xDD, 0xEF, 0xE3),
    "◐": RGBColor(0xFB, 0xEC, 0xD2),
    "–": RGBColor(0xF3, 0xF3, 0xF3),
}
SYM_FG = {
    "●": RGBColor(0x1E, 0x7A, 0x4C),
    "◐": RGBColor(0xA0, 0x74, 0x18),
    "–": RGBColor(0xB8, 0xB8, 0xB8),
}


def _set_cell(cell, text, size, color, bold=False, italic=False, fill=None, align=PP_ALIGN.LEFT,
              font=FONT_BODY, valign=MSO_ANCHOR.MIDDLE, line_spacing=0.95):
    cell.margin_left = Inches(0.04)
    cell.margin_right = Inches(0.04)
    cell.margin_top = Inches(0.02)
    cell.margin_bottom = Inches(0.02)
    cell.vertical_anchor = valign
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = WHITE
    tf = cell.text_frame
    tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, (line_text, line_size, line_color, line_bold, line_italic) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line_text
        r.font.size = Pt(line_size)
        r.font.bold = line_bold
        r.font.italic = line_italic
        r.font.color.rgb = line_color
        r.font.name = font


def add_capability_matrix_slide(prs, tenet, idx, total, finish_slide_fn):
    data = MATRIX[tenet]
    columns = data["columns"]
    rows = data["rows"]
    color = TENET_COLORS.get(tenet, TEAL)

    slide = blank_slide(prs)
    add_header(slide, f"Capability Matrix {idx} of {total}", tenet, accent=color)
    add_text(slide, 0.55, 1.42, 12.2, 0.28,
              f"Every specific control found for this tenet across {len(columns)} referenced tools - "
              "rows are capabilities, columns are tools, with a cost/efficiency best pick per row.",
              size=10, color=TEXT_MUTED, font=FONT_BODY)

    label_w = 3.05
    n_cols = len(columns)
    other_w = (12.2 - label_w) / n_cols
    n_rows = len(rows) + 1
    header_h = 0.34
    body_h = 0.47 if len(rows) <= 9 else 0.43

    table_top = 1.78
    graphic_frame = slide.shapes.add_table(n_rows, n_cols + 1, Inches(0.55), Inches(table_top),
                                            Inches(12.2), Inches(header_h + body_h * len(rows)))
    table = graphic_frame.table
    table.columns[0].width = Inches(label_w)
    for i in range(n_cols):
        table.columns[i + 1].width = Inches(other_w)
    table.rows[0].height = Inches(header_h)
    for i in range(1, n_rows):
        table.rows[i].height = Inches(body_h)

    _set_cell(table.cell(0, 0), [("CAPABILITY  /  BEST PICK", 8, WHITE, True, False)], 8, WHITE,
              fill=NAVY, align=PP_ALIGN.LEFT)
    for j, col in enumerate(columns):
        _set_cell(table.cell(0, j + 1), [(col["code"], 8.3, WHITE, True, False)], 8.3, WHITE,
                  fill=NAVY, align=PP_ALIGN.CENTER)

    for i, row in enumerate(rows):
        r = i + 1
        label_cell = table.cell(r, 0)
        lines = [(row["aspect"], 7.6, TEXT_DARK, True, False)]
        if row["best"]:
            lines.append((row["best"], 6.6, TEXT_MUTED, False, True))
        _set_cell(label_cell, lines, 8, TEXT_DARK, fill=CARD_BG if i % 2 == 0 else WHITE)
        for j, col in enumerate(columns):
            sym = row["cells"].get(col["code"], "–")
            cell = table.cell(r, j + 1)
            _set_cell(cell, [(sym, 12, SYM_FG[sym], True, False)], 12, SYM_FG[sym],
                      fill=SYM_FILL[sym], align=PP_ALIGN.CENTER)

    legend_y = table_top + header_h + body_h * len(rows) + 0.12
    add_text(slide, 0.55, legend_y, 12.2, 0.22,
              "●  native / dedicated capability      ◐  partial or delegated to another tool      –  not found",
              size=9, color=TEXT_DARK, font=FONT_BODY)
    code_list = "   ".join(f"{c['code']}={c['name']}" for c in columns)
    add_text(slide, 0.55, legend_y + 0.24, 12.2, 0.42, code_list, size=7.6, color=TEXT_MUTED,
              font=FONT_BODY, line_spacing=1.15)

    finish_slide_fn(slide, f"Capability Matrix - {tenet}")


def add_all_capability_matrices(prs, finish_slide_fn):
    total = len(TENET_ORDER)
    for i, tenet in enumerate(TENET_ORDER, start=1):
        add_capability_matrix_slide(prs, tenet, i, total, finish_slide_fn)
