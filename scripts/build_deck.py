# -*- coding: utf-8 -*-
"""
Builds AFNI_Responsible_AI_Framework.pptx end to end.
"""
import json
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "helpers"))

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from build_pptx import (
    new_prs, blank_slide, add_rect, add_rounded, add_text, add_bullets,
    add_footer, add_header, pill, tenet_pills, badge, add_table,
    NAVY, NAVY_DARK, TEAL, BG_LIGHT, CARD_BG, WHITE, TEXT_DARK, TEXT_MUTED,
    TEXT_SOFT_ON_NAVY, AMBER, GREEN, RED_SOFT, LINE_GREY, TENET_COLORS,
    TENET_ORDER, ROLE_COLORS, FONT_HEAD, FONT_BODY, SLIDE_W_IN, SLIDE_H_IN,
    OUT_PATH, SYNTHESIS_PATH,
)
from repo_slide_content import REPO_SLIDES

PAGE = [0]  # mutable page counter


def next_page():
    PAGE[0] += 1
    return PAGE[0]


def finish_slide(slide, section):
    add_footer(slide, next_page(), section)


# --------------------------------------------------------------- SLIDE 1 ----
def slide_title(prs):
    slide = blank_slide(prs, NAVY)
    add_rect(slide, 0, 0, SLIDE_W_IN, SLIDE_H_IN, NAVY)
    add_rect(slide, 0, 2.55, SLIDE_W_IN, 0.045, TEAL)
    add_text(slide, 0.9, 1.55, 11.5, 0.5, "RESPONSIBLE AI GOVERNANCE", size=16, color=TEAL,
              bold=True, font=FONT_HEAD)
    add_text(slide, 0.85, 2.05, 11.6, 1.5,
              "A Responsible AI Framework for AFNI", size=40, color=WHITE, bold=True, font=FONT_HEAD)
    add_text(slide, 0.9, 2.95, 11.4, 0.7,
              "Reviewing 23 open-source tools and designing one unified toolkit,\n"
              "tenet by tenet, for every AI application AFNI builds.",
              size=15, color=TEXT_SOFT_ON_NAVY, font=FONT_BODY, line_spacing=1.3)
    add_rect(slide, 0.9, 6.35, 0.5, 0.04, TEAL)
    add_text(slide, 0.9, 6.5, 6, 0.35, "Prepared by Sai Muthiki", size=13, color=WHITE, bold=True, font=FONT_HEAD)
    add_text(slide, 0.9, 6.82, 8, 0.35, "For Kiran Devkar & AFNI AI Governance", size=12,
              color=TEXT_SOFT_ON_NAVY, font=FONT_BODY)
    add_text(slide, 10.6, 6.5, 2.2, 0.6, "AFNI", size=22, color=WHITE, bold=True,
              align=PP_ALIGN.RIGHT, font=FONT_HEAD)
    next_page()
    return slide


# --------------------------------------------------------------- SLIDE 2 ----
def slide_agenda(prs):
    slide = blank_slide(prs)
    add_header(slide, "Agenda", "What This Deck Covers")
    items = [
        ("01", "Why This Matters", "The business case and the regulatory backdrop driving this work"),
        ("02", "The 7 Tenets", "What \u201cResponsible AI\u201d means in practical terms at AFNI"),
        ("03", "Framework Architecture", "Where guardrail checks live, and development vs. testing"),
        ("04", "23 Repository Deep-Dives", "One slide per tool - features, limits, cost, and fit"),
        ("05", "Master Checklist", "Every concrete check found, tagged to a tenet"),
        ("06", "Tenet-by-Tenet Recommendations", "The best combination of tools per tenet"),
        ("07", "Feasibility & Roadmap", "What to adopt, when, and in what order"),
    ]
    y = 1.65
    for num, title, desc in items:
        add_text(slide, 0.7, y, 0.9, 0.5, num, size=22, color=TEAL, bold=True, font=FONT_HEAD)
        add_text(slide, 1.55, y + 0.02, 4.4, 0.4, title, size=15, color=NAVY, bold=True, font=FONT_HEAD)
        add_text(slide, 5.9, y + 0.02, 6.6, 0.5, desc, size=12, color=TEXT_MUTED, font=FONT_BODY)
        if num != "07":
            add_rect(slide, 0.7, y + 0.62, 12.0, 0.012, LINE_GREY)
        y += 0.73
    finish_slide(slide, "Agenda")


# --------------------------------------------------------------- SLIDE 3 ----
def slide_why(prs):
    slide = blank_slide(prs)
    add_header(slide, "Business Context", "Why This Matters to AFNI")
    add_bullets(slide, 0.6, 1.65, 6.6, 5.2, [
        "AFNI is moving to AI-native development. Every new AI app should follow one shared safety standard, not a different approach each time.",
        "Two reasons this pays off: (1) it builds AI that AFNI and its people can actually trust, and (2) it helps AFNI pass client security reviews before demoing or shipping AI products.",
        "Regulators are catching up fast. The EU AI Act's toughest rules apply from August 2026. The US NIST AI framework and ISO 42001 are becoming the standard questions clients ask vendors.",
        "AFNI's cloud is Azure-first, but not locked in. The goal is to pick the best tool for each job, whether it's open-source or a paid cloud service.",
        "This deck reviews 23 real, hands-on tools - not just theory - to turn \u201cwe should be responsible\u201d into a concrete, buildable plan.",
    ], size=14.5, line_spacing=1.15, space_after=14)
    card = add_rounded(slide, 7.6, 1.65, 5.2, 5.15, NAVY, radius=0.05)
    add_text(slide, 7.9, 1.95, 4.6, 0.4, "KEY DATES", size=12, color=TEAL, bold=True, font=FONT_HEAD)
    dates = [
        ("Aug 2026", "EU AI Act high-risk obligations take effect"),
        ("Ongoing", "NIST AI RMF used as the de-facto US reference"),
        ("Ongoing", "ISO/IEC 42001 turns AI ethics into a certifiable standard"),
        ("Client asks", "\u201cWhich framework do you follow?\u201d is now a standard question"),
    ]
    y = 2.5
    for date, desc in dates:
        add_text(slide, 7.9, y, 4.6, 0.3, date, size=13, color=WHITE, bold=True, font=FONT_HEAD)
        add_text(slide, 7.9, y + 0.32, 4.6, 0.55, desc, size=11.5, color=TEXT_SOFT_ON_NAVY,
                  font=FONT_BODY, line_spacing=1.1)
        y += 1.05
    finish_slide(slide, "Why This Matters")


# --------------------------------------------------------------- SLIDE 4 ----
TENET_BLURBS = {
    "Privacy": "Protecting personal data - names, SSNs, medical records - so it never leaks into an AI answer.",
    "Security": "Defending against jailbreaks, prompt injection, and other attacks that trick an AI into misbehaving.",
    "Fairness & Bias": "Making sure AI decisions treat every group of people fairly, with no hidden discrimination.",
    "Explainability & Transparency": "Being able to show why an AI made a decision, in plain language.",
    "Profanity / Content Safety": "Blocking toxic, hateful, or otherwise unsafe language in and out of the AI.",
    "Hallucination / Reliability": "Catching made-up facts and making sure answers are grounded in real information.",
    "Accountability": "Clear ownership, audit trails, and logging for every AI decision and system change.",
}


def slide_tenets(prs):
    slide = blank_slide(prs)
    add_header(slide, "The Foundation", "What \u201cResponsible AI\u201d Means at AFNI")
    add_text(slide, 0.6, 1.42, 11.8, 0.4,
              "Different frameworks use different names, but they converge on the same seven ideas:",
              size=13, color=TEXT_MUTED, font=FONT_BODY)
    cols = 4
    card_w, card_h = 2.9, 2.05
    gap_x, gap_y = 0.18, 0.2
    start_x, start_y = 0.55, 1.95
    for i, tenet in enumerate(TENET_ORDER):
        col = i % cols
        row = i // cols
        x = start_x + col * (card_w + gap_x)
        y = start_y + row * (card_h + gap_y)
        color = TENET_COLORS[tenet]
        add_rounded(slide, x, y, card_w, card_h, CARD_BG, radius=0.06, line=True)
        add_rect(slide, x, y, card_w, 0.09, color)
        add_text(slide, x + 0.18, y + 0.22, card_w - 0.36, 0.5, tenet, size=13, color=NAVY,
                  bold=True, font=FONT_HEAD, line_spacing=1.0)
        add_text(slide, x + 0.18, y + 0.85, card_w - 0.36, card_h - 1.0, TENET_BLURBS[tenet],
                  size=10.5, color=TEXT_MUTED, font=FONT_BODY, line_spacing=1.12)
    add_text(slide, 0.55, 6.55, 11.8, 0.5,
              "NIST calls this \u201ctrustworthy AI\u201d: valid, reliable, safe, secure, accountable, transparent, "
              "explainable, private, and fair.", size=11.5, color=TEXT_MUTED, italic=True, font=FONT_BODY)
    finish_slide(slide, "The 7 Tenets")


# --------------------------------------------------------------- SLIDE 5 ----
def slide_architecture(prs):
    slide = blank_slide(prs)
    add_header(slide, "Framework Design", "Two Layers: Build the Guardrail, Test the Guardrail")
    add_text(slide, 0.6, 1.42, 11.8, 0.55,
              "A responsible AI toolkit needs both halves working together - tools that enforce safety live, "
              "and tools that attack/grade it before and after launch.",
              size=13, color=TEXT_MUTED, font=FONT_BODY, line_spacing=1.15)

    dev_x, test_x = 0.6, 6.75
    col_w = 5.9
    top_y = 2.15
    add_rounded(slide, dev_x, top_y, col_w, 0.55, TEAL, radius=0.15)
    add_text(slide, dev_x, top_y, col_w, 0.55, "GUARDRAIL DEVELOPMENT  (build the defense)", size=13,
              color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=FONT_HEAD)
    add_bullets(slide, dev_x, top_y + 0.75, col_w, 3.6, [
        "Runs live, in front of and behind every AI call",
        "Examples: PII redaction, jailbreak filters, content-safety checks",
        "Tools: NeMo Guardrails, Guardrails AI, LLM Guard, Infosys Toolkit, Presidio-style checks, SHAP, Fairlearn/AIF360",
    ], size=12.5, bullet_color=TEAL, space_after=10)

    add_rounded(slide, test_x, top_y, col_w, 0.55, RED_SOFT, radius=0.15)
    add_text(slide, test_x, top_y, col_w, 0.55, "VULNERABILITY / RED-TEAM TESTING  (attack the defense)", size=13,
              color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=FONT_HEAD)
    add_bullets(slide, test_x, top_y + 0.75, col_w, 3.6, [
        "Runs offline, before release and on a schedule, in CI/CD",
        "Examples: jailbreak fuzzing, red-team attacks, benchmark scoring",
        "Tools: PyRIT, garak, FuzzyAI, LLMFuzzer, JCB, DeepTeam, promptfoo, giskard, deepeval, evals",
    ], size=12.5, bullet_color=RED_SOFT, space_after=10)

    add_rect(slide, 0.6, 6.05, 12.1, 0.02, LINE_GREY)
    add_text(slide, 0.6, 6.2, 12.1, 0.85,
              "How they connect: red-team and evaluation findings become the pass/fail gate in CI/CD. "
              "A guardrail change only ships once it survives the attack suite - so testing tools continuously "
              "sharpen the development-layer tools, not the other way around.",
              size=12.5, color=TEXT_DARK, font=FONT_BODY, line_spacing=1.2)
    finish_slide(slide, "Framework Architecture")


# --------------------------------------------------------------- DIVIDER ----
def slide_divider(prs, kicker, title, subtitle):
    slide = blank_slide(prs, NAVY)
    add_rect(slide, 0, 0, SLIDE_W_IN, SLIDE_H_IN, NAVY)
    add_rect(slide, 0.9, 3.15, 0.6, 0.05, TEAL)
    add_text(slide, 0.9, 2.5, 11, 0.4, kicker.upper(), size=14, color=TEAL, bold=True, font=FONT_HEAD)
    add_text(slide, 0.85, 3.3, 11.6, 1.1, title, size=34, color=WHITE, bold=True, font=FONT_HEAD)
    add_text(slide, 0.9, 4.35, 10.8, 0.8, subtitle, size=14, color=TEXT_SOFT_ON_NAVY, font=FONT_BODY,
              line_spacing=1.25)
    next_page()
    return slide


# --------------------------------------------------------- REPO SLIDE ----
def slide_repo(prs, r, idx, total):
    slide = blank_slide(prs)
    primary_color = TENET_COLORS.get(r["tenets"][0], TEAL)
    add_header(slide, f"Repository {idx} of {total}", r["display_name"], accent=primary_color)

    left_x, left_w = 0.55, 7.2
    right_x, right_w = 8.0, 4.75

    # Left column
    add_text(slide, left_x, 1.48, left_w, 1.05, r["summary"], size=11.5, color=TEXT_DARK, italic=True,
              font=FONT_BODY, line_spacing=1.15)
    y = 2.62
    add_text(slide, left_x, y, left_w, 0.28, "KEY FEATURES", size=11, color=TEAL, bold=True, font=FONT_HEAD)
    add_bullets(slide, left_x, y + 0.3, left_w, 2.15, r["features"], size=10.2, bullet_color=TEAL,
                space_after=5, line_spacing=1.0)
    y2 = 5.17
    add_text(slide, left_x, y2, left_w, 0.28, "LIMITATIONS TO WATCH", size=11, color=AMBER, bold=True,
              font=FONT_HEAD)
    add_bullets(slide, left_x, y2 + 0.3, left_w, 1.6, r["limitations"], size=9.8, bullet_color=AMBER,
                space_after=4, line_spacing=1.0)

    # Right column - tenet pills + role/layer + facts card + fit note
    y = 1.5
    tenet_end_y = tenet_pills(slide, right_x, y, r["tenets"], max_w=right_w, size=9.5)
    y = tenet_end_y + 0.1
    role_color = ROLE_COLORS.get(r["role"], TEAL)
    role_label = {"Guardrail Development": "DEVELOPMENT", "Vulnerability / Red-Team Testing": "RED-TEAM TESTING",
                  "Both": "DEVELOPMENT + TESTING"}[r["role"]]
    badge(slide, right_x, y, role_label, role_color, w=right_w, h=0.32, size=10)
    y += 0.48

    card_h = 1.92
    add_rounded(slide, right_x, y, right_w, card_h, CARD_BG, radius=0.06, line=True)
    vendor_short = r["vendor"] if len(r["vendor"]) <= 34 else r["vendor"][:33].rsplit(" ", 1)[0] + "…"
    facts = [
        ("Layer type", r["layer"]), ("Tier", r["tier"]),
        ("Cost model", r["cost"]), ("Vendor", vendor_short),
        ("Integration effort", r["effort"]), ("License", r["license"]),
    ]
    col_w = (right_w - 0.16 * 2 - 0.1) / 2
    fy0 = y + 0.13
    for i, (label, val) in enumerate(facts):
        col = i % 2
        row = i // 2
        fx = right_x + 0.16 + col * (col_w + 0.1)
        fy = fy0 + row * 0.56
        add_text(slide, fx, fy, col_w, 0.2, label.upper(), size=8.3, color=TEXT_MUTED,
                  bold=True, font=FONT_HEAD)
        add_text(slide, fx, fy + 0.19, col_w, 0.34, val, size=9.8, color=NAVY,
                  bold=True, font=FONT_BODY, wrap=True, line_spacing=0.95)
    y += card_h + 0.13

    fit_h = 7.03 - y
    add_rounded(slide, right_x, y, right_w, fit_h, NAVY, radius=0.06)
    add_text(slide, right_x + 0.15, y + 0.09, right_w - 0.3, 0.22, "AFNI FIT", size=9.5, color=TEAL, bold=True,
              font=FONT_HEAD)
    fit_lines = max(2, min(5, int(len(r["fit"]) / 62) + 1))
    fit_text_h = fit_lines * 0.155 + 0.05
    add_text(slide, right_x + 0.15, y + 0.32, right_w - 0.3, fit_text_h, r["fit"], size=9.3,
              color=WHITE, font=FONT_BODY, line_spacing=1.08)
    bvb_y = y + 0.32 + fit_text_h + 0.1
    add_text(slide, right_x + 0.15, bvb_y, right_w - 0.3, 0.2, "BUILD VS. BUY", size=9.5, color=TEAL,
              bold=True, font=FONT_HEAD)
    add_text(slide, right_x + 0.15, bvb_y + 0.23, right_w - 0.3, fit_h - (bvb_y - y) - 0.23,
              r["build_replicate"], size=9.3, color=WHITE, font=FONT_BODY, line_spacing=1.08)

    finish_slide(slide, r["display_name"])


# ------------------------------------------------------------- BUILD -----
def build():
    prs = new_prs()
    slide_title(prs)
    slide_agenda(prs)
    slide_why(prs)
    slide_tenets(prs)
    slide_architecture(prs)

    slide_divider(prs, "Section", "23 Repository Deep-Dives",
                  "One slide per tool: what it does, its real features and limits, cost, "
                  "effort to integrate, and where it fits into AFNI's toolkit.")
    total = len(REPO_SLIDES)
    for i, r in enumerate(REPO_SLIDES, start=1):
        slide_repo(prs, r, i, total)

    synthesis = None
    if os.path.exists(SYNTHESIS_PATH):
        with open(SYNTHESIS_PATH, encoding="utf-8") as f:
            synthesis = json.load(f)

    if synthesis:
        import build_deck_synthesis as bds
        bds.add_synthesis_slides(prs, synthesis, slide_divider, finish_slide, next_page)
    else:
        print("WARNING: synthesis file not found yet - synthesis slides skipped in this build.")

    prs.save(OUT_PATH)
    print(f"Saved {OUT_PATH} with {next_page()-1} slides tracked (PAGE counter).")
    print(f"Actual slide count: {len(prs.slides.__iter__.__self__._sldIdLst)}")


if __name__ == "__main__":
    build()
